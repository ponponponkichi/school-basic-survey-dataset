from __future__ import annotations

from io import BytesIO
import zipfile

import pandas as pd
from openpyxl.styles import Font, PatternFill


SOURCE_NOTE = (
    "出典：大学改革支援・学位授与機構「大学基本情報」 "
    "https://portal.niad.ac.jp/ptrt/table.html を加工して作成"
)


def _sheet_name(form_code: str) -> str:
    return form_code[:31]


def _format_sheet(worksheet) -> None:
    header_fill = PatternFill(patternType="solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.iter_cols(min_row=1, max_row=min(worksheet.max_row, 200)):
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 10), 32)


def build_excel_workbook(
    frames: dict[str, pd.DataFrame],
    manifest: pd.DataFrame,
    column_dictionary: pd.DataFrame,
) -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for form_code, dataframe in frames.items():
            dataframe.to_excel(writer, sheet_name=_sheet_name(form_code), index=False)
        manifest.to_excel(writer, sheet_name="取得情報", index=False)
        column_dictionary.to_excel(writer, sheet_name="列定義", index=False)
        note = pd.DataFrame({"説明": [SOURCE_NOTE]})
        note.to_excel(writer, sheet_name="利用上の注意", index=False)
        for worksheet in writer.book.worksheets:
            _format_sheet(worksheet)
    return buffer.getvalue()


def build_csv_zip(
    frames: dict[str, pd.DataFrame],
    manifest: pd.DataFrame,
    column_dictionary: pd.DataFrame,
) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        for form_code, dataframe in frames.items():
            archive.writestr(
                f"{form_code}.csv",
                dataframe.to_csv(index=False).encode("utf-8-sig"),
            )
        archive.writestr(
            "取得情報.csv",
            manifest.to_csv(index=False).encode("utf-8-sig"),
        )
        archive.writestr(
            "列定義.csv",
            column_dictionary.to_csv(index=False).encode("utf-8-sig"),
        )
        archive.writestr("README.txt", (SOURCE_NOTE + "\n").encode("utf-8"))
    return buffer.getvalue()

